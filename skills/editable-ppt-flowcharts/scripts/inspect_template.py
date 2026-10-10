"""Read template styling statistics without copying slides or media."""

import json
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shapes(items):
    for shape in items:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes(shape.shapes)


def inspect(path):
    prs = Presentation(path)
    fonts, sizes, colors = Counter(), Counter(), Counter()
    for slide in prs.slides:
        for shape in shapes(slide.shapes):
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts[run.font.name] += 1
                        if run.font.size:
                            sizes[run.font.size.pt] += 1
                        if run.font.color.type == MSO_COLOR_TYPE.RGB:
                            colors[str(run.font.color.rgb)] += 1
                        for ea in run._r.xpath("./a:rPr/a:ea"):
                            fonts[ea.get("typeface", "")] += 1
            if hasattr(shape, "fill") and shape.fill.type == MSO_FILL_TYPE.SOLID:
                if shape.fill.fore_color.type == MSO_COLOR_TYPE.RGB:
                    colors[str(shape.fill.fore_color.rgb)] += 1
            if hasattr(shape, "line") and shape.line.color.type == MSO_COLOR_TYPE.RGB:
                colors[str(shape.line.color.rgb)] += 1
    return {"canvas": [prs.slide_width / 914400, prs.slide_height / 914400],
            "slides": len(prs.slides), "explicit_font_counts": dict(fonts.most_common()),
            "explicit_font_size_counts": dict(sizes.most_common()),
            "explicit_rgb_counts": dict(colors.most_common()),
            "note": "Explicit values only; inspect rendering for inherited theme fonts/colors."}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: python inspect_template.py reference.pptx")
    print(json.dumps(inspect(Path(sys.argv[1]).expanduser()), ensure_ascii=False, indent=2))
