"""Validate native editable shapes separately from PowerPoint rendering evidence."""
import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.util import Inches


def normalized(text):
    return " ".join(str(text).split())


def validate(pptx, spec_path=None, report_path=None):
    prs = Presentation(pptx)
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8-sig")) if spec_path else None
    errors, texts, expected_bounds = [], [], set()
    allowed = {getattr(MSO_SHAPE_TYPE, name) for name in ("AUTO_SHAPE", "TEXT_BOX", "LINE")}
    pictures = 0
    for index, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.shape_type not in allowed:
                errors.append(f"Slide {index}: non-native shape {shape.name} ({shape.shape_type})")
            pictures += int(shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE))
            if (shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 2
                    or shape.top + shape.height > prs.slide_height + 2):
                errors.append(f"Slide {index}: shape outside canvas: {shape.name}")
            if shape.has_text_frame and shape.text.strip():
                if spec is not None:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if not run.text.strip():
                                continue
                            if run.font.name != spec.get("font", "Arial"):
                                errors.append(f"Slide {index}: font mismatch in {shape.name}: {run.font.name}")
                            ea = run._r.xpath("./a:rPr/a:ea")
                            ea_font = ea[0].get("typeface") if ea else None
                            if ea_font != spec.get("east_asian_font", "Microsoft YaHei"):
                                errors.append(f"Slide {index}: East Asian font mismatch in {shape.name}: {ea_font}")
                hidden = shape._element.xpath('.//*[local-name()="cNvPr"][@hidden="1"]')
                if not hidden:
                    slide_texts.append(shape.text)
                    expected_bounds.add((index, shape.name))
        if not slide_texts:
            errors.append(f"Slide {index}: no nonempty editable text")
        texts.append(normalized(" ".join(slide_texts)))
    if not prs.slides:
        errors.append("Presentation has no slides")
    edge_status = "not_requested"
    if spec_path:
        if len(spec["slides"]) != len(prs.slides):
            errors.append("Slide count differs from specification")
        for label in spec.get("required_visible_labels", []):
            if normalized(label) not in " ".join(texts):
                errors.append(f"Missing required label: {label}")
        edge_status = "passed"
        for index, (slide, slide_spec) in enumerate(zip(prs.slides, spec["slides"]), 1):
            named = {shape.name: shape for shape in slide.shapes}
            labels = [slide_spec.get(key, "") for key in ("title", "subtitle", "footer")]
            labels += slide_spec.get("required_visible_labels", [])
            for label in filter(None, labels):
                if normalized(label) not in texts[index - 1]:
                    errors.append(f"Slide {index}: missing required label: {label}")
            for node in slide_spec.get("nodes", []):
                shape = named.get("node:" + node["id"])
                if shape is None or not shape.has_text_frame:
                    errors.append(f"Slide {index}: missing editable node: {node['id']}")
                    continue
                for label in [node.get("title", "")] + node.get("body", []):
                    if label and normalized(label) not in normalized(shape.text):
                        errors.append(f"Slide {index}: node {node['id']} missing text: {label}")
            edge_names = [name for name in named if name.startswith("edge:")]
            if slide_spec.get("edges") and not edge_names:
                edge_status = "unavailable_names"
            expected_edges = set()
            for edge_index, edge in enumerate(slide_spec.get("edges", []), 1):
                points = edge["points"]
                is_block = edge.get("style", "line") == "block"
                segments = [(points[0], points[-1])] if is_block else list(zip(points, points[1:]))
                for segment_index, (start, end) in enumerate(segments, 1):
                    name = f"edge:{edge['from']}->{edge['to']}:{edge_index}:{segment_index}"
                    expected_edges.add(name)
                    shape = named.get(name)
                    if shape is None:
                        errors.append(f"Slide {index}: missing edge segment: {name}")
                        continue
                    if is_block:
                        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or shape.auto_shape_type != MSO_SHAPE.RIGHT_ARROW:
                            errors.append(f"Slide {index}: expected RIGHT_ARROW: {name}")
                            continue
                        actual = (shape.left, shape.top, shape.width, shape.height)
                        expected = tuple(Inches(v) for v in (start[0], start[1] - 0.125, end[0] - start[0], 0.25))
                        transforms = shape._element.xpath("./p:spPr/a:xfrm")
                        if any(t.get("flipH") == "1" for t in transforms):
                            errors.append(f"Slide {index}: reversed block arrow: {name}")
                    else:
                        geometry = shape._element.xpath("./p:spPr/a:prstGeom")
                        if (shape.shape_type != MSO_SHAPE_TYPE.LINE or len(geometry) != 1
                                or geometry[0].get("prst") != "line"):
                            errors.append(f"Slide {index}: expected straight connector: {name}")
                            continue
                        actual = (shape.begin_x, shape.begin_y, shape.end_x, shape.end_y)
                        expected = tuple(Inches(v) for v in (*start, *end))
                        heads = shape._element.xpath(".//a:ln/a:headEnd")
                        tails = shape._element.xpath(".//a:ln/a:tailEnd")
                        if any(item.get("type", "none") != "none" for item in heads):
                            errors.append(f"Slide {index}: reverse headEnd arrow: {name}")
                        needs_arrow = segment_index == len(segments) and edge.get("arrow", True)
                        if needs_arrow:
                            if len(tails) != 1 or tails[0].get("type") != "triangle":
                                errors.append(f"Slide {index}: missing or incorrect tailEnd arrow: {name}")
                        elif any(item.get("type", "none") != "none" for item in tails):
                            errors.append(f"Slide {index}: unexpected tailEnd arrow: {name}")
                    if shape.rotation != 0 or any(abs(a - b) > 2 for a, b in zip(actual, expected)):
                        errors.append(f"Slide {index}: edge geometry differs from specification: {name}")
            for name in sorted(set(edge_names) - expected_edges):
                errors.append(f"Slide {index}: unexpected edge segment: {name}")
        if errors and edge_status == "passed":
            edge_status = "failed"
    rendered = {"status": "not_performed", "errors": []}
    if report_path:
        rendered_errors = rendered["errors"]
        report = json.loads(Path(report_path).read_text(encoding="utf-8-sig"))
        summary = report.get("summary", {})
        if report.get("pptx_sha256") != hashlib.sha256(Path(pptx).read_bytes()).hexdigest():
            rendered_errors.append("Render report does not identify the current PPTX SHA-256")
        if report.get("status") != "passed" or report.get("error"):
            rendered_errors.append(report.get("error") or "PowerPoint report is not passed")
        if report.get("overflow") != [] or summary.get("overflow_count") != 0:
            rendered_errors.append("Text overflow or missing overflow results")
        checked = {(item.get("slide"), item.get("textboundshape")) for item in report.get("text_bounds", [])}
        if checked != expected_bounds or summary.get("text_shapes_checked") != len(expected_bounds):
            rendered_errors.append("Rendered text-bound coverage differs from visible editable text")
        for item in report.get("text_bounds", []):
            bounds, shape_bounds = item.get("bounds", []), item.get("shape_bounds", [])
            if len(bounds) != 4 or len(shape_bounds) != 4:
                rendered_errors.append("Missing actual text/shape bounds")
            elif (bounds[0] < shape_bounds[0] - 1 or bounds[1] < shape_bounds[1] - 1
                  or bounds[0] + bounds[2] > shape_bounds[0] + shape_bounds[2] + 1
                  or bounds[1] + bounds[3] > shape_bounds[1] + shape_bounds[3] + 1):
                rendered_errors.append(f"Actual text bounds exceed shape: {item.get('textboundshape')}")
        if summary.get("rendered_slides") != len(prs.slides) or len(report.get("images", [])) != len(prs.slides):
            rendered_errors.append("Slide rendering is incomplete")
        for image in report.get("images", []):
            image_path = Path(image.get("path", ""))
            if not image_path.is_file() or image_path.read_bytes()[:8] != b"\x89PNG\r\n\x1a\n":
                rendered_errors.append(f"Rendered PNG missing or invalid: {image_path}")
        rendered["status"] = "failed" if rendered_errors else "passed"
    return {"pptx": str(Path(pptx).resolve()), "slides": len(prs.slides),
            "structural": {"status": "failed" if errors else "passed", "errors": errors,
                           "picture_shapes": pictures, "spec_edge_check": edge_status}, "rendered": rendered}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx")
    parser.add_argument("--spec")
    parser.add_argument("--render-report")
    args = parser.parse_args()
    try:
        result = validate(args.pptx, args.spec, args.render_report)
    except Exception as exc:
        result = {"structural": {"status": "failed", "errors": [str(exc)]},
                  "rendered": {"status": "not_performed", "errors": []}}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(int(result["structural"]["status"] != "passed" or result["rendered"]["status"] == "failed"))
