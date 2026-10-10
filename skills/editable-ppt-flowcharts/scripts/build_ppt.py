"""Draw an editable PowerPoint flowchart from the documented layout JSON."""

import json
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

PALETTE = {
    "blue": ("2962A4", "DCEBFC"), "green": ("2E8B57", "DCF5E6"),
    "orange": ("E67E22", "FFF0DC"), "purple": ("8E44AD", "F0E1FA"),
    "gray": ("B4B4B4", "F0F0F0"),
}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def validate(spec):
    require(isinstance(spec, dict), "Spec must be a JSON object")
    canvas = spec.get("canvas", [13.333333, 7.5])
    require(isinstance(canvas, list) and len(canvas) == 2 and
            all(number(v) and v > 0 for v in canvas), "Invalid canvas")
    width, height = canvas
    require(width >= 2 and height >= 1.4, "Canvas is too small for the header and page number")
    for key in ("font", "east_asian_font"):
        require(isinstance(spec.get(key, "Arial"), str) and spec.get(key, "Arial").strip(), f"Invalid {key}")
    slides = spec.get("slides")
    require(isinstance(slides, list) and slides, "Spec requires nonempty slides")
    for index, slide in enumerate(slides, 1):
        prefix = f"Slide {index}: "
        require(isinstance(slide, dict) and isinstance(slide.get("title"), str) and
                slide["title"].strip(), prefix + "title required")
        for key in ("subtitle", "notes", "footer"):
            require(isinstance(slide.get(key, ""), str), prefix + f"invalid {key}")
        nodes = slide.get("nodes")
        require(isinstance(nodes, list) and nodes, prefix + "nodes required")
        ids = set()
        for node in nodes:
            require(isinstance(node, dict), prefix + "invalid node")
            node_id = node.get("id")
            require(isinstance(node_id, str) and node_id.strip() and node_id not in ids,
                    prefix + "missing or duplicate node ID")
            ids.add(node_id)
            rect = node.get("rect")
            require(isinstance(rect, list) and len(rect) == 4 and all(number(v) for v in rect),
                    prefix + f"{node_id}: invalid rect")
            x, y, w, h = rect
            require(x >= 0 and y >= 0 and w > 0 and h > 0 and
                    x + w <= width + 1e-6 and y + h <= height + 1e-6,
                    prefix + f"{node_id}: rect outside canvas or nonpositive size")
            require(node.get("kind", "box") in ("box", "text"), prefix + "invalid node kind")
            require(isinstance(node.get("title", ""), str) and
                    isinstance(node.get("body", []), list) and
                    all(isinstance(v, str) for v in node.get("body", [])), prefix + "invalid node text")
            require(node.get("accent", "blue") in PALETTE, prefix + "invalid node accent")
            require(node.get("align", "left") in ALIGN, prefix + "invalid alignment")
            for key in ("font_size", "body_font_size"):
                require(key not in node or (number(node[key]) and node[key] > 0), prefix + f"invalid {key}")
            require(isinstance(node.get("bold", True), bool), prefix + "invalid bold")
        edges = slide.get("edges", [])
        require(isinstance(edges, list), prefix + "invalid edges")
        for edge in edges:
            require(isinstance(edge, dict) and isinstance(edge.get("from"), str) and
                    isinstance(edge.get("to"), str) and edge["from"] in ids and edge["to"] in ids,
                    prefix + "edge references missing node")
            points = edge.get("points")
            require(isinstance(points, list) and len(points) >= 2, prefix + "edge needs two or more points")
            for point in points:
                require(isinstance(point, list) and len(point) == 2 and all(number(v) for v in point) and
                        0 <= point[0] <= width and 0 <= point[1] <= height, prefix + "invalid/outside edge point")
            require(all(a != b for a, b in zip(points, points[1:])), prefix + "zero-length edge segment")
            require(edge.get("style", "line") in ("line", "block"), prefix + "invalid edge style")
            require(edge.get("accent", "blue") in PALETTE, prefix + "invalid edge accent")
            for key in ("arrow", "dashed"):
                require(isinstance(edge.get(key, key == "arrow"), bool), prefix + f"invalid {key}")
            if edge.get("style") == "block":
                require(len(points) == 2 and points[0][1] == points[1][1] and
                        points[1][0] > points[0][0], prefix + "block requires two horizontal points increasing in x")
                require(0.125 <= points[0][1] <= height - 0.125, prefix + "block arrow outside canvas")
                require(edge.get("arrow", True) and not edge.get("dashed", False),
                        prefix + "block arrows require arrow=true and dashed=false")
    return width, height


def no_shadow(shape):
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    style = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style is not None:
        effect = style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef")
        if effect is not None:
            effect.set("idx", "0")


def paragraph(frame, text, size, bold, ink, align, font, ea_font, first=False, gap=0):
    p = frame.paragraphs[0] if first else frame.add_paragraph()
    p.alignment = align
    p.space_before, p.space_after, p.line_spacing = Pt(0), Pt(gap), 1.0
    run = p.add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(ink)
    ea = OxmlElement("a:ea")
    ea.set("typeface", ea_font)
    run._r.get_or_add_rPr().append(ea)


def textbox(slide, rect, text, size, bold, ink, font, ea_font, name, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(*(Inches(v) for v in rect))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.01)
    if align == PP_ALIGN.RIGHT:
        frame.margin_right = Inches(0.06)
    frame.margin_top = frame.margin_bottom = 0
    paragraph(frame, text, size, bold, ink, align, font, ea_font, first=True)
    no_shadow(shape)
    return shape


def draw_node(slide, node, font, ea_font):
    accent, fill = PALETTE[node.get("accent", "blue")]
    is_box = node.get("kind", "box") == "box"
    rect = tuple(Inches(v) for v in node["rect"])
    if is_box:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *rect)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.color.rgb, shape.line.width = RGBColor.from_string(accent), Pt(1.2)
        shape.adjustments[0] = 0.12
    else:
        shape = slide.shapes.add_textbox(*rect)
    shape.name = "node:" + node["id"]
    no_shadow(shape)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.10 if is_box else 0.01)
    frame.margin_top = frame.margin_bottom = Inches(0.05 if is_box else 0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE if is_box else MSO_ANCHOR.TOP
    align = ALIGN[node.get("align", "center" if is_box else "left")]
    ink = "333333" if node.get("accent") == "gray" else accent
    paragraph(frame, node.get("title", ""), node.get("font_size", 18 if is_box else 19),
              True if is_box else node.get("bold", True), ink, align, font, ea_font,
              first=True, gap=5 if node.get("body") else 0)
    for text in node.get("body", []):
        paragraph(frame, text, node.get("body_font_size", 16), False, "333333",
                  align, font, ea_font, gap=1)


def draw_edge(slide, edge, edge_index):
    ink = RGBColor.from_string(PALETTE[edge.get("accent", "blue")][0])
    points = edge["points"]
    name = f"edge:{edge['from']}->{edge['to']}:{edge_index}"
    if edge.get("style", "line") == "block":
        x, y = points[0]
        shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y - 0.125),
                                       Inches(points[1][0] - x), Inches(0.25))
        shape.name = name + ":1"
        shape.fill.solid()
        shape.fill.fore_color.rgb = ink
        shape.line.fill.background()
        no_shadow(shape)
        return
    for index, (start, end) in enumerate(zip(points, points[1:])):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                          *(Inches(v) for v in (*start, *end)))
        shape.name = f"{name}:{index + 1}"
        shape.line.color.rgb, shape.line.width = ink, Pt(1.7)
        ln = shape._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
        if edge.get("dashed", False):
            dash = OxmlElement("a:prstDash")
            dash.set("val", "dash")
            ln.append(dash)
        if index == len(points) - 2 and edge.get("arrow", True):
            arrow = OxmlElement("a:tailEnd")
            for key, value in {"type": "triangle", "w": "sm", "len": "sm"}.items():
                arrow.set(key, value)
            ln.append(arrow)
        no_shadow(shape)


def build(spec, output):
    width, height = validate(spec)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    font, ea_font = spec.get("font", "Arial"), spec.get("east_asian_font", "Microsoft YaHei")
    count = len(spec["slides"])
    for index, data in enumerate(spec["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        textbox(slide, (0.5, 0.25, width - 1, 0.58), data["title"], 32, True,
                "2962A4", font, ea_font, "Slide title")
        if data.get("subtitle"):
            textbox(slide, (0.5, 0.88, width - 1, 0.35), data["subtitle"], 16.5, False,
                    "333333", font, ea_font, "Slide subtitle")
        for edge_index, edge in enumerate(data.get("edges", []), 1):
            draw_edge(slide, edge, edge_index)
        for node in data["nodes"]:
            draw_node(slide, node, font, ea_font)
        if data.get("footer"):
            textbox(slide, (0.5, height - 0.40, width - 2, 0.27), data["footer"], 14, False,
                    "737373", font, ea_font, "Slide footer")
        textbox(slide, (width - 1.18, height - 0.22, 0.65, 0.20), f"{index} / {count}",
                9.5, False, "737373", font, ea_font, "Page number", PP_ALIGN.RIGHT)
        slide.notes_slide.notes_text_frame.text = data.get("notes", "")
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.title = spec["slides"][0]["title"]
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(json.dumps({"output": str(output), "slides": count,
                      "native_shapes_per_slide": [len(s.shapes) for s in prs.slides]}))


def main():
    require(len(sys.argv) == 3, "Usage: python build_ppt.py layout.json output.pptx")
    source, output = (Path(v).expanduser().resolve() for v in sys.argv[1:])
    require(source != output, "Output must not overwrite the input JSON")
    require(output.suffix.lower() == ".pptx", "Output must have a .pptx extension")
    build(json.loads(source.read_text(encoding="utf-8-sig")), output)


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError) as error:
        print(f"Error: {error}", file=sys.stderr)
        sys.exit(1)
